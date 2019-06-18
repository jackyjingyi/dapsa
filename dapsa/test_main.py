import pandas as pd
from pdfminer.pdfparser import PDFParser
from pdfminer.pdfdocument import PDFDocument
from pdfminer.pdfpage import PDFTextExtractionNotAllowed
from pdfminer.pdfinterp import PDFResourceManager
from pdfminer.pdfinterp import PDFPageInterpreter
from pdfminer.pdfdevice import PDFDevice,PDFTextDevice
from pdfminer.layout import LAParams,LTChar,LTTextLine,LTText,LTTextBox,LTLine,LTRect,LTImage,LTTextBoxHorizontal
from pdfminer.converter import PDFPageAggregator
from pdfminer.pdfpage import PDFPage
from dapsa.test_class_v2 import Point,Cluster,Cell,Row,aCell
from pdfminer.utils import Plane
from core_function.preparation.pdf_settings import COUNTRY_LIST
import pdfminer
import re
import copy
import openpyxl
import time
import datetime
import nltk

import os

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "settings")

from django.core.wsgi import get_wsgi_application
application = get_wsgi_application()

from pdfstage.models import Document
from ReviewStage.models import SpecLevel

def approximation(xc, x0, width=3, pts=0.2):
    return abs(xc - x0) < width * (1 + pts)


def checkpointsvaliadation(point_list, textboxlist):
    # remove all points in a textbox e.g a html address with bottom line
    valid_points = []
    unvalid_points = []
    for point in point_list:
        unvalid = False
        for t in textboxlist:
            if all([point.x >= t.bbox[0], point.x <= t.bbox[2],
                    point.y >= t.bbox[1], point.y <= t.bbox[3]]):
                unvalid = True
        if unvalid:
            unvalid_points.append(point)
        else:
            valid_points.append(point)

    print(len(unvalid_points))
    return valid_points


def group_cluster(cleaned_points):
    cluster_container = []
    for i in range(len(cleaned_points)):
        if len(cluster_container) > 0:
            found = False
            for cluster in cluster_container:
                if cluster.is_acceptable(cleaned_points[i]) and cleaned_points[
                    i] not in cluster.value:  # due to its a set , doesnt matter if duplicates
                    cluster.fill_(cleaned_points[i])
                    found = True
            if not found:
                new_cluster = Cluster()
                new_cluster.fill_(cleaned_points[i])
                cluster_container.append(new_cluster)
        else:
            new_cluster = Cluster()
            new_cluster.fill_(cleaned_points[i])
            cluster_container.append(new_cluster)
    return cluster_container


# this stage is grouping clusters based on their Y value
# 1. graping all stages from list : stage of y values, grouping them
# this step is base structure for building cells
def clsmanager(cluster_container):
    temp = cluster_container
    holder = {}
    while len(temp) > 0:
        max_value = round(max([i.key[1] for i in temp]))
        current = temp.pop(0)
        if len(holder) > 0:
            check = [approximation(k, current.key[1]) for k in holder.keys()]

            if any(check):
                # key is here
                # get the current_key index mapping back to dict
                current_key = list(holder.keys())[check.index(True)]
                holder[current_key] = holder[current_key] + [current]
            else:
                # if not shows in the dict need to build one
                holder[max_value] = [current]
        else:
            holder[max_value] = [current]
    return holder


# now we get holder on hand  which contains a lots of clusters
# next step is to build a cell by all the clusters
# a cell is what we calculate the center of 4 clusters then, its close to the top left cluster


def cls2cells(holder, top_limit):
    holder_key_list = list(holder.keys())

    def sort_cluster(c):
        return c.key[0]

    for h in holder.values():
        h.sort(key=sort_cluster)
    cell_container = []
    for i in range(len(holder_key_list) - 1):  # [762,729,705,680,675,571,536,432,340,259,212,154,96,61]
        checking_list = copy.deepcopy(holder[holder_key_list[i]])
        while len(checking_list) > 1:
            checking_item = checking_list.pop(0)
            if holder_key_list[i] < top_limit:  # [675,571,536,432,340,259,212,154,96,61]
                cls_index = [i.same_cluster(checking_item) for i in holder[holder_key_list[i]]].index(True)
                a_cell = Cell(holder[holder_key_list[i]][cls_index])
                for j in range(cls_index + 1, len(holder[holder_key_list[i]])):
                    a_cell.rtc_update(holder[holder_key_list[i]][j])
                for j in range(cls_index, len(holder[holder_key_list[i + 1]])):
                    a_cell.lbc_update(holder[holder_key_list[i + 1]][j])
                    a_cell.rbc_update(holder[holder_key_list[i + 1]][j])
                cell_container.append(a_cell)
    return cell_container


def cleancells(cell_container):
    # cells settle down, oh yeah
    cleaned_cell_container = []
    for cell in cell_container:
        if not any([c == None for c in cell.cluster_collection.values()]):
            cell.set_center()
            cleaned_cell_container.append(cell)
    return cleaned_cell_container


def puttext2cells(cleaned_cell_container, textbox):
    for text in textbox:
        for r in cleaned_cell_container:
            if r.textbelongto(text):
                r.set_text(text)
            else:
                # print('#######\n' ,text)
                pass


# put cells into a row collection and column collection
# more importantly is update row and column id for each cell 4/25
# now let's get started jingyi from 4/26

def cells2rows(cleaned_cell_container):
    # c is cleaned_cell , r = []
    def sortkeyforrow(r):
        return r.centre_y_coordinates

    row_container = []
    while len(cleaned_cell_container) > 0:
        current = cleaned_cell_container.pop()
        if len(row_container) == 0:
            a_row = Row()
            a_row.set_row(current)
            row_container.append(a_row)
        else:
            found = False
            for row in row_container:
                if row.isbelong(current):
                    row.fill_cell(current)
                    found = True
            if not found:
                a_row = Row()
                a_row.set_row(current)
                row_container.append(a_row)
    for r in row_container:
        r.sort_row()
        for c in r.cells_list:
            pass
            # print(c.identification['centre_coordinates'],c.get_text())
    # print([r.centre_y_coordinates for r in row_container])
    row_container.sort(key=sortkeyforrow, reverse=True)
    return row_container


def read_output(loc, file, countries_list):
    token_grage = {}
    df = pd.read_excel(loc + file, sheet_name='output', index_col=0)
    df_original = pd.read_excel(loc + file, sheet_name='Sheet1', index_col=0)
    # select Fail culumn with x
    # not na and not '-
    df1 = df.loc[(df['Fail'].notna()) & (
                (df['Fail'] == 'X \n') | (df['Fail'] == 'x \n') | (df['Fail'] == 'Fail \n'))]  # , df['Fail'] =='x \n')]
    # print('Total {} section failed in report'.format(len(df1.index)))
    df1.to_excel('C:/Users/yanjingy/PycharmProjects/JanBI/PDFmining/Failed_row_from_report/' + file + '.xlsx')
    # print(df1)
    # col[0] is test_item
    target = df1[df1.columns[0]]
    # store the info of failed section extracted from test report, inspection references
    # df1.to_excel('.\extract_data_stroage\df_fina2.xlsx', sheet_name= 'output',header = False)
    # get token and store it into grage , comment s is last col .values[0] means the first value in numpy list ,here we only
    # get one value
    for item in target:
        token_grage[item] = ([word.lower() for word in nltk.word_tokenize(item) if word.isalpha()],
                             df1[df1[df1.columns[0]] == item][df1.columns[-1]].values[0])
    # what's on protocol
    # the first col is not na and only grab the first column
    dftest = df_original[df_original[df_original.columns[0]].notna()][df_original.columns[0]]

    # dftest_withmethod = df_original[df_original['Test_method'].notna()]['Test_method']  # not used yet
    xindex = []
    for i in range(len(list(dftest))):
        current = [word.lower() for word in nltk.word_tokenize(list(dftest)[i]) if word.isalpha()]

        for token in token_grage.values():

            if current == token[0]:
                xindex.append((dftest.iloc[i], token[1]))
                # print('**********')
                # print(dftest.iloc[i])
                # print(len(xindex))
                # print(token[0])
                # print(current)
                break
            else:
                pass
    # get the row which matched
    df_final = df_original.loc[df_original[df_original.columns[0]].isin([token[0] for token in xindex])]
    insert_dfn = len(df_final[df_final.columns[0]]) * ['DFN']
    insert_issue = len(df_final[df_final.columns[0]]) * [
        'Test (non-measurement test) documented on test report, failed']
    df_final.insert(5, 'scenario', insert_dfn)
    df_final.insert(6, 'comments', [token[1] for token in xindex])
    df_final.insert(7, 'issue', insert_issue)

    # write it to an excel
    # df_final.to_excel('./final_outputs/df_final.xlsx', sheet_name= 'output')  #########)
    df_final = df_final[df_final[df_final.columns[1]].isin(countries_list)]
    return df_final


"""
def read_output(loc,file,countries_list):
    token_grage = {}
    df = pd.read_excel(loc+file, sheet_name='output',index_col=0)
    df_original = pd.read_excel(loc+file,sheet_name='Sheet1',index_col=0)
    # select Fail culumn with x
    # not na and not '-
    df1 = df.loc[(df['Results'].notna())&(df['Results'].str.contains('Fail'))]
    print('Total {} section failed in report'.format(len(df1.index)))
    print(df1)
    # col[0] is test_item
    target = df1[df1.columns[0]]
    # store the info of failed section extracted from test report, inspection references
    # df1.to_excel('.\extract_data_stroage\df_fina2.xlsx', sheet_name= 'output',header = False)
    # get token and store it into grage , comment s is last col .values[0] means the first value in numpy list ,here we only
    # get one value
    for item in target:
        token_grage[item] = ([word.lower() for word in nltk.word_tokenize(item) if word.isalpha()],
                             df1[df1[df1.columns[0]]==item][df1.columns[-1]].values[0])
    # what's on protocol
    # the first col is not na and only grab the first column
    dftest = df_original[df_original[df_original.columns[0]].notna()][df_original.columns[0]]
    print(type(dftest),dftest)
    print(list(dftest)[1])
    # dftest_withmethod = df_original[df_original['Test_method'].notna()]['Test_method']  # not used yet
    xindex = []
    for i in range(len(dftest)):
        current =[word.lower() for word in nltk.word_tokenize(list(dftest)[i]) if word.isalpha()]

        for token in token_grage.values():

            if current == token[0]:
                xindex.append((dftest.iloc[i],token[1]))
                print(len(xindex))
                print(token[0])
                print(current)
                break
            else:
                pass
    # get the row which matched
    df_final = df_original.loc[df_original[df_original.columns[0]].isin([token[0] for token in xindex])]
    insert_dfn = len(df_final[df_final.columns[0]]) *['DFN']
    insert_issue = len(df_final[df_final.columns[0]]) *['Test (non-measurement test) documented on test report, failed']
    df_final.insert(5,'scenario',insert_dfn)
    df_final.insert(6,'comments',[token[1] for token in xindex])
    df_final.insert(7,'issue',insert_issue)
    print('###################')
    print(df_final)
    # write it to an excel
    #df_final.to_excel('./final_outputs/df_final.xlsx', sheet_name= 'output')  #########
    #print(df_final.loc[10])
    df_final = df_final[df_final[df_final.columns[1]].isin(countries_list)]
    return df_final"""


def build_df(length, login, audit_date, protocol, mfp):
    login_list = [login] * length
    audit_date_list = [audit_date] * length
    protocol_list = [protocol] * length
    mfp_list = [mfp] * length
    df = pd.DataFrame({
        'Login': login_list,
        'Date': audit_date_list,
        'Protocol': protocol_list,
        'MFR Part': mfp_list
    })
    print(df)
    return df


if __name__ == '__main__':
    start_timestamp = time.time()
    # file inputs pagenumber , [points] , [textbox]
    # open a pdf file
    a = Document.objects.get(pk=3).document.path
    fp = open(a, 'rb')
    target_pages = [i for i in range(8,21)]
    # create a pdf parser object associated with the file object
    parser = PDFParser(fp)

    # create a pdf document object that stores the document structure
    # supply the password for initialization
    document = PDFDocument(parser)
    if not document.is_extractable:
        raise PDFTextExtractionNotAllowed

    # Create a pdf resource manager object that stores shared resources
    rsrcmgr = PDFResourceManager()
    # Create a PDF device object
    device = PDFDevice(rsrcmgr)
    textdevice = PDFTextDevice(device)

    # Process each page contained in the document

    laparams = LAParams(line_margin=0.01, char_margin=0.1)
    pageagg = PDFPageAggregator(rsrcmgr, laparams=laparams)
    new_interpreter = PDFPageInterpreter(rsrcmgr, pageagg)
    pagedict = {}
    # {page_index in pdf file : gageid}  unnecessary time consumption refine next round
    for pageid, page in enumerate(PDFPage.get_pages(fp), start=1):
        pagedict[pageid] = page.pageid

    print(pagedict)

    raw_data_storage = {}


    def getpageinfo(page, interpreter=new_interpreter):
        textbox_list = []
        interpreter.process_page(page)
        layout = pageagg.get_result()
        Point.setpage_Points(pageid=page.pageid)
        point_count = 0
        mid_count = 0
        for element in layout:
            if isinstance(element, LTChar):
                print(element.fontname)
            elif isinstance(element, pdfminer.layout.LTCurve):
                print(element)
                point_left_bottom = Point(x=element.bbox[0], y=element.bbox[1], pageid=page.pageid, id=point_count)
                point_right_top = Point(element.bbox[2], element.bbox[3], pageid=page.pageid, id=point_count + 1)
                middle_point = Point(x=(element.bbox[0] + element.bbox[2]) / 2,
                                     y=(element.bbox[1] + element.bbox[3]) / 2, pageid=page.pageid, id=mid_count,
                                     ischild=True)
                point_left_bottom.addTopagePoints(pageid=page.pageid)
                point_right_top.addTopagePoints(pageid=page.pageid)
                point_left_bottom.assignchild(middle_point)
                point_right_top.assignchild(middle_point)
                middle_point.addTopagePoints(pageid=page.pageid)
                # print(point_left_bottom.__repr__())
                # print(point_right_top.__repr__())
                # print(point_left_bottom.children)
                # print(point_right_top.children)
                # print(middle_point.parents)
                point_count += 2
                mid_count += 1
            elif isinstance(element, LTTextBox):
                print(element)
                textbox_list.append(element)
            elif isinstance(element, pdfminer.layout.LTLine):
                print('this is lttline {}'.format(element))
        return Point.pagePoints[page.pageid], textbox_list


    page_warehouse = {}
    for page in PDFPage.get_pages(fp):
        for tp in target_pages:
            if page.pageid == pagedict[tp]:
                page_output = getpageinfo(page)
                page_warehouse[(tp, page.pageid)] = page_output  # (points , textbox)


    def extracttable(pageid, textbox, top_limit=800):
        print('the len is {}'.format(len(Point.pagePoints[pageid])))
        clean_points = Point.cleanPoints(pageid=pageid, textboxlist=textbox)
        Cluster.group_cluster(pageid=pageid, cleaned_points=clean_points)
        Cluster.sortpageCluster(pageid=pageid)
        Cluster.assignid(pageid=pageid)

        for c in Cluster.pageSortedClusters[pageid]:
            c.get_middle()
            print(c.id, c.s1_matrix, c.key)
        print(Cluster.pageSortedClusters)
        # checking point 1
        Cluster.locat_roots(pageid=pageid)
        Cluster.do_cell(pageid=pageid)
        for c in list(aCell.pageCells[pageid].values()):
            for j in c:
                j()
                print((j.left_boundary, j.lower_boundary, j.right_boundary, j.upper_boundary, j.width, j.height,
                       j.centre))

        holder = clsmanager([c for c in Cluster.pageSortedClusters[pageid] if len(c.value) > 3])
        cell_container = cls2cells(holder, top_limit=top_limit)
        cleaned_cells = cleancells(cell_container)
        puttext2cells(cleaned_cells, textbox)
        row_container = cells2rows(cleaned_cells)
        aCell.AssignText(pageid=pageid, textlist=textbox)
        aCell.check_text(pageid=pageid)
        print(aCell.pageCells)
        return row_container


    output_wharehouse = {}
    for key in page_warehouse.keys():
        r = extracttable(pageid=key[1], textbox=page_warehouse[key][1])
        output_wharehouse[key] = r


    def pagesort(x):
        return x[0]


    final_output_temp_container = [output_wharehouse[key] for key in
                                   sorted(list(output_wharehouse.keys()), key=pagesort)]
    final_output_list = []
    for l in final_output_temp_container:
        final_output_list += l
    output_data1 = pd.DataFrame([[c.get_text() for c in r.cells_list] for r in final_output_list])

    # output_data1 = pd.DataFrame([[c.get_text() for c in r.cells_list] for r in final_output_list])
    # output_data1.to_excel('check.xlsx')
    output_data = pd.DataFrame([[c.get_text() for c in r.cells_list] for r in final_output_list],
                                columns= ['TESTItem', 'Country', 'Requirement', 'Pass', 'Fail', 'COMMENTS'])

    protocol = pd.read_excel('C:/Users/yanjingy/PycharmProjects/JanBI/PDFmining/protocols/complete/AMZ-034-WW_One Time Use Tableware_V2.0_20180629.xlsx',
                             sheet_name='Sheet1')
    # print(protocol)
    with pd.ExcelWriter('C:/Users/yanjingy/PycharmProjects/JanBI/PDFmining/extract_data_stroage/33-X-testi' + '-output' + '8-20-6' + '.xlsx') as writer:
        output_data.to_excel(writer, sheet_name='output')
        protocol.to_excel(writer, sheet_name='Sheet1', index=0)

    #############################
    # analysis the extrated data
    out_loc = 'C:/Users/yanjingy/PycharmProjects/JanBI/PDFmining/final_outputs/'
    out_file = 'HPB_target.xlsx'
    # fm = config.confile['Fail-mark']

    cl = COUNTRY_LIST
    media = read_output(loc='C:/Users/yanjingy/PycharmProjects/JanBI/PDFmining/extract_data_stroage/',
                        file='33-X-testi' + '-output' + '8-20-6' + '.xlsx',
                        countries_list=cl)
    length = len(media.index)

    wb = openpyxl.load_workbook(out_loc + out_file)
    ##$$
    ws = wb['Spec Level']
    login = ws['A2'].value
    audit_date = datetime.datetime.today().strftime('%m/%d/%Y')
    protocol = 'AMZ-034-WW_One Time Use Tableware_V2.0_20180629'
    mfp = '33-X-testing-(8-20-6`)-<05-07-2019>'
    speck_level_df = build_df(length, login=login, audit_date=audit_date, protocol=protocol, mfp=mfp)
    # insert column into the new dataframe
    speck_level_df.insert(4, 'Section of the protocol', list(media[media.columns[4]]))
    speck_level_df.insert(5, 'TEST DESCRIPTION', list(media[media.columns[0]]))
    speck_level_df.insert(6, 'Test method listed on protocol', list(media[media.columns[2]]))
    speck_level_df.insert(7, 'Issue found with the test/test method (reason for fail)', list(media[media.columns[-1]]))
    end_timestamp = time.time()
    timerange = round(end_timestamp - start_timestamp, 3)
    speck_level_df.insert(8, 'Safety & Regulatory', [str(timerange) + ' secs'] * length)
    speck_level_df.insert(9, 'Additional notes regarding "Issue found with the test/test method (reason for fail)" ',
                          list(media[media.columns[-2]]))
    speck_level_df.insert(10, 'Remark', list(media[media.columns[1]]))

    if len(list(speck_level_df.index)) > 0:
        for i in range(len(list(speck_level_df.index))):
            m = list(speck_level_df.loc[i])
            print(m)
            a = SpecLevel(
                auditor=m[0],
                protocol=m[2],
                mfp=m[3],
                section=m[4],
                description=m[5],
                method=m[6],
                issue=m[7],
                sr=m[8],
                addition=m[9],
                remark=m[10],
                audit_date=datetime.datetime.today())
            a.save()


    print('time spent on this program {} secs'.format(round(timerange, 2)))